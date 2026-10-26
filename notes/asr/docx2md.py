import re
from docx import Document

SRC = r"D:/AI/0、公司内部项目/项目归档工具/notes/029业务档案归档管理办法-2026-4-10.docx"
DST = r"D:/AI/0、公司内部项目/项目归档工具/notes/029业务档案归档管理办法-2026-4-10.md"

doc = Document(SRC)
paras = [p.text.rstrip() for p in doc.paragraphs]

out = []
title_done = False
prev_numbered = False
chapter_no = 0

for idx, raw in enumerate(paras):
    t = raw.strip()

    if not t:
        prev_numbered = False
        continue

    # 封面标题区（前三行）
    if idx == 0:
        out.append(f"# {t}\n")
        continue
    if idx == 1:
        out.append(f"## {t}\n")
        continue
    if idx == 2:
        out.append(f"> {t}\n")
        out.append("---\n")
        continue

    # 章
    if re.match(r"^第[一二三四五六七八九十]+章", t):
        chapter_no += 1
        out.append(f"\n## {t}\n")
        prev_numbered = False
        continue

    # 条
    m = re.match(r"^(第[一二三四五六七八九十]+条)\s*(.*)$", t)
    if m:
        out.append(f"\n### {m.group(1)}　{m.group(2)}\n")
        prev_numbered = False
        continue

    # 编号条目 1、xxx
    m = re.match(r"^(\d+)、\s*(.*)$", t)
    if m:
        out.append(f"{m.group(1)}. {m.group(2)}")
        prev_numbered = True
        continue

    # 注：及附注
    if t.startswith("注：") or re.match(r"^注[：:]", t):
        out.append(f"\n**{t}**\n")
        prev_numbered = False
        continue

    # 编号条目的续行（如“含特许经营方案……”“含委托人提供的……”）
    if prev_numbered and ("含" in t[:4] or raw.startswith("    ") or raw.startswith("\t")):
        # 挂到上一条后面
        for j in range(len(out) - 1, -1, -1):
            if re.match(r"^\d+\. ", out[j]):
                out[j] = out[j] + f"<br>　　{t}"
                break
        continue

    # 附件标题等
    if t.startswith("附件"):
        out.append(f"\n---\n\n### {t}\n")
        prev_numbered = False
        continue

    out.append(t)
    prev_numbered = False

# 表格（附件1 交接表）
if doc.tables:
    t = doc.tables[0]
    rows = t.rows
    out.append("\n")
    out.append("| 序号 | 文件名称 | 存档日期 | 备注 |")
    out.append("|---|---|---|---|")
    # 第一行是“项目名称”合并行，单独处理
    header_cells = [c.text.strip() for c in rows[0].cells]
    out.insert(len(out) - 2, "")
    for r in rows[1:]:
        cells = [c.text.strip().replace("\n", " ") for c in r.cells]
        out.append("| " + " | ".join(cells) + " |")

text = "\n".join(out)
text = re.sub(r"\n{3,}", "\n\n", text)

# 表格前插入“项目名称”行说明
text = text.replace(
    "| 序号 | 文件名称 | 存档日期 | 备注 |",
    "**项目名称：**\n\n| 序号 | 文件名称 | 存档日期 | 备注 |",
)

# 末尾补充签署行
text += "\n归档人员：　　　　　　　　　　　　接收人员：\n"

open(DST, "w", encoding="utf-8").write(text)
print("saved", DST)
print("chars", len(text))
