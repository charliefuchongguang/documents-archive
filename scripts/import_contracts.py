#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
导入公司贴花合同台账（默认 sheet：公司贴花合同四）。

用法：
  python scripts/import_contracts.py --dry-run
  python scripts/import_contracts.py --apply

合同性质只保留原文，不做归一化；历史合同默认不是大合同，不影响项目关联。
"""
import argparse
import csv
import datetime as dt
import json
import os
import re
import sqlite3
import sys
import uuid

try:
    import openpyxl
except ImportError:
    print('缺少 openpyxl，请先安装：pip install openpyxl', file=sys.stderr)
    sys.exit(1)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(ROOT, 'notes', '公司贴花合同 2012.xlsx')
DEFAULT_DB = os.path.join(ROOT, 'src', 'server', 'database', 'archive.db')
DEFAULT_REPORT = os.path.join(ROOT, 'notes', '合同台账导入异常报告.csv')
SHEET_NAME = '公司贴花合同四'
LEGACY_SOURCE = 'legacy'
ZERO_WIDTH_RE = re.compile('[\u200b-\u200d\ufeff]')


def clean_text(value):
    if value is None:
        return ''
    s = str(value)
    s = ZERO_WIDTH_RE.sub('', s)
    return s.replace('\u3000', ' ').strip()


def add_issue(issues, row_no, field, raw, message):
    issues.append({
        'row': row_no,
        'field': field,
        'raw': '' if raw is None else str(raw),
        'message': message,
    })


def parse_number(value):
    if value is None or value == '':
        return None, False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), False
    s = clean_text(value).replace(',', '')
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    if not m:
        return None, True
    try:
        return float(m.group(0)), False
    except ValueError:
        return None, True


def parse_int(value):
    n, invalid = parse_number(value)
    if n is None:
        return None, invalid
    return int(n), invalid


def days_in_month(year, month):
    import calendar
    return calendar.monthrange(year, month)[1]


def parse_partial_date(value):
    """返回 (normalized, start, end, issue)。支持 YYYY / YYYY-MM / YYYY-MM-DD。"""
    if value is None or value == '':
        return None, None, None, None
    if isinstance(value, dt.datetime):
        d = value.date()
        s = d.isoformat()
        return s, s, s, None
    if isinstance(value, dt.date):
        s = value.isoformat()
        return s, s, s, None

    raw = clean_text(value)
    if not raw or raw in ('无', '-', '—'):
        return None, None, None, '合同日期为空或无'

    s = raw.strip('（）()【】[] ')
    s = re.sub(r'[至—~～]', '/', s)
    if '/' in s:
        s = [p.strip() for p in s.split('/') if p.strip()][-1]
    s = s.replace('年', '-').replace('月', '-').replace('日', '')
    s = re.sub(r'-+', '-', s).strip('- ')
    if not s:
        return None, None, None, '合同日期格式无法解析'

    parts = s.split('-')
    # 兼容 20020-6-20 这类 5 位年份笔误
    if parts and len(parts[0]) == 5 and parts[0].startswith('20'):
        parts[0] = parts[0][:2] + parts[0][3:]
    s = '-'.join(parts)

    m = re.match(r'^(\d{4})(?:-(\d{1,2})(?:-(\d{1,2}))?)?$', s)
    if not m:
        return None, None, None, '合同日期格式无法解析'
    year = int(m.group(1))
    month = int(m.group(2)) if m.group(2) else None
    day = int(m.group(3)) if m.group(3) else None
    if year < 1900 or year > 2100:
        return None, None, None, '合同日期年份异常'
    if month is not None and (month < 1 or month > 12):
        return None, None, None, '合同日期月份异常'
    if day is not None and (day < 1 or day > days_in_month(year, month)):
        return None, None, None, '合同日期天数异常'

    if month is None:
        normalized = f'{year:04d}'
        start = f'{year:04d}-01-01'
        end = f'{year:04d}-12-31'
    elif day is None:
        normalized = f'{year:04d}-{month:02d}'
        start = f'{year:04d}-{month:02d}-01'
        end = f'{year:04d}-{month:02d}-{days_in_month(year, month):02d}'
    else:
        normalized = f'{year:04d}-{month:02d}-{day:02d}'
        start = normalized
        end = normalized
    return normalized, start, end, None


def parse_full_date(value):
    if value is None or value == '':
        return None, None
    if isinstance(value, dt.datetime):
        return value.date().isoformat(), None
    s = clean_text(value)
    if not s or s in ('无', '-', '—'):
        return None, '存档日期为空或无'
    m = re.match(r'^(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})', s)
    if not m:
        return None, '存档日期格式无法解析'
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return dt.date(y, mo, d).isoformat(), None
    except ValueError:
        return None, '存档日期无效'


def parse_credit_codes(value):
    raw = clean_text(value)
    if not raw:
        return '', []
    parts = re.split(r'[/、,，;\n]+', raw)
    codes = []
    for p in parts:
        p = clean_text(p).strip()
        if p and p not in codes:
            codes.append(p)
    return raw, codes


def infer_payment(name):
    s = clean_text(name)
    if not s:
        return False
    # （付）、(付2025)、名称末尾“付”等；不把“支付方式”误判为付款合同
    return bool(re.search(r'(?:[（(]\s*付[^）)]*[）)]|付\s*$|付\s*20\d{2})', s))


def read_rows(path, sheet_name):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise RuntimeError(f'找不到工作表：{sheet_name}')
    ws = wb[sheet_name]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not values or all(v is None or clean_text(v) == '' for v in values):
            continue
        row = list(values) + [None] * max(0, 17 - len(values))
        rows.append({
            'excel_row': excel_row,
            'values': row,
            'headers': headers,
        })
    wb.close()
    return rows


def build_records(path, sheet_name):
    rows = read_rows(path, sheet_name)
    valid_rows = [r for r in rows if clean_text(r['values'][2])]
    int_seqs = []
    for r in valid_rows:
        seq, _ = parse_int(r['values'][1])
        if seq and seq > 0:
            int_seqs.append(seq)
    max_seq = max(int_seqs) if int_seqs else 0
    used_seqs = set(int_seqs)
    next_seq = max_seq

    records = []
    report_rows = []
    issue_counts = {}
    nature_values = {}
    payment_count = 0
    consolidated_count = 0
    duplicate_names = {}

    for r in valid_rows:
        values = r['values']
        excel_row = r['excel_row']
        issues = []
        legacy_extra = {'issues': [], 'raw': {}}

        name = clean_text(values[2])
        if not name:
            continue
        duplicate_names[name] = duplicate_names.get(name, 0) + 1

        seq, seq_invalid = parse_int(values[1])
        if not seq:
            if excel_row == 220 and 219 not in used_seqs:
                seq = 219
            else:
                next_seq += 1
                seq = next_seq
            used_seqs.add(seq)
            add_issue(issues, excel_row, '序号', values[1], f'序号缺失，已补为 {seq}')
        elif seq_invalid:
            add_issue(issues, excel_row, '序号', values[1], f'序号非数字，已按 {seq} 处理')

        counterparty = clean_text(values[3])
        total_investment, total_invalid = parse_number(values[4])
        if total_invalid:
            add_issue(issues, excel_row, '总投资额（万元）', values[4], '金额格式异常')
            legacy_extra['raw']['total_investment'] = str(values[4])

        nature_raw = clean_text(values[5])
        nature_values[nature_raw] = nature_values.get(nature_raw, 0) + 1
        if not nature_raw:
            add_issue(issues, excel_row, '合同性质', values[5], '合同性质为空，保留原文为空')
        elif re.match(r'^\d+$', nature_raw):
            add_issue(issues, excel_row, '合同性质', values[5], '合同性质为数字，保留原文，待后续人工确认')

        contract_content = clean_text(values[6])
        project_manager = clean_text(values[7])
        tech_manager = clean_text(values[8])

        if isinstance(values[9], dt.datetime):
            contract_date_raw = values[9].date().isoformat()
        elif isinstance(values[9], dt.date):
            contract_date_raw = values[9].isoformat()
        else:
            contract_date_raw = clean_text(values[9])
        contract_date, contract_start, contract_end, date_issue = parse_partial_date(values[9])
        if date_issue:
            add_issue(issues, excel_row, '合同日期', values[9], date_issue)
        elif isinstance(values[9], str):
            raw_date_text = clean_text(values[9])
            if re.match(r'^\d{5}-', raw_date_text):
                add_issue(issues, excel_row, '合同日期', values[9], '年份疑似 5 位笔误，已按修正年份处理')
            if re.search(r'[/至—~～]', raw_date_text):
                add_issue(issues, excel_row, '合同日期', values[9], '包含多个日期，已取最后一个')

        agency_fee, fee_invalid = parse_number(values[10])
        if fee_invalid:
            add_issue(issues, excel_row, '代理费金额（万元）', values[10], '代理费非数字，已置空')
            legacy_extra['raw']['agency_fee'] = str(values[10])

        # 早期列错位：应贴花列出现日期时，不写入金额字段
        due_raw = values[11]
        paid_raw = values[12]
        if isinstance(due_raw, dt.datetime):
            stamp_tax_due = None
            add_issue(issues, excel_row, '应贴花金额（元）', due_raw, '该列出现日期，疑似历史列错位，金额置空')
            legacy_extra['raw']['stamp_tax_due'] = due_raw.isoformat()
        else:
            stamp_tax_due, due_invalid = parse_number(due_raw)
            if due_invalid:
                add_issue(issues, excel_row, '应贴花金额（元）', due_raw, '金额格式异常')
                legacy_extra['raw']['stamp_tax_due'] = str(due_raw)

        stamp_tax_paid, paid_invalid = parse_number(paid_raw)
        if paid_invalid:
            add_issue(issues, excel_row, '已贴花金额（元）', paid_raw, '金额格式异常')
            legacy_extra['raw']['stamp_tax_paid'] = str(paid_raw)
        if isinstance(due_raw, dt.datetime) and paid_raw not in (None, ''):
            legacy_extra['raw']['legacy_paid_or_copies'] = str(paid_raw)

        copies, copies_invalid = parse_int(values[13])
        if copies_invalid:
            add_issue(issues, excel_row, '份数', values[13], '份数非数字，已置空')
            legacy_extra['raw']['copies'] = str(values[13])
        elif copies is not None and copies > 100:
            add_issue(issues, excel_row, '份数', values[13], f'份数异常（{copies}），疑似历史贴花金额，已置空')
            legacy_extra['raw']['copies'] = str(values[13])
            copies = None

        archive_date, archive_issue = parse_full_date(values[14])
        if archive_issue:
            add_issue(issues, excel_row, '存档日期', values[14], archive_issue)

        tax_raw = clean_text(values[15])
        is_consolidated_tax = 1 if tax_raw else 0
        if tax_raw:
            consolidated_count += 1
            legacy_extra['raw']['汇总纳税'] = tax_raw

        credit_code_raw, credit_codes = parse_credit_codes(values[16])
        if values[16] is not None and clean_text(values[16]) != str(values[16]).strip():
            # 零宽字符/首尾空格已经清理，不需要重复报错
            pass

        is_payment = infer_payment(name)
        if is_payment:
            payment_count += 1

        if issues:
            for issue in issues:
                issue_counts[issue['field']] = issue_counts.get(issue['field'], 0) + 1
                report_rows.append({
                    '行号': excel_row,
                    '序号': seq,
                    '项目名称': name,
                    '问题字段': issue['field'],
                    '原始值': issue['raw'],
                    '建议处理': issue['message'],
                })
            legacy_extra['issues'] = [f"{it['field']}: {it['message']}" for it in issues]

        records.append({
            'id': str(uuid.uuid4()),
            'seq_no': seq,
            'name': name,
            'counterparty': counterparty,
            'total_investment': total_investment,
            'nature_raw': nature_raw,
            'contract_content': contract_content,
            'project_manager': project_manager,
            'tech_manager': tech_manager,
            'contract_date_raw': contract_date_raw,
            'contract_date': contract_start or contract_date,
            'contract_date_end': contract_end,
            'agency_fee': agency_fee,
            'stamp_tax_due': stamp_tax_due,
            'stamp_tax_paid': stamp_tax_paid,
            'copies': copies,
            'archive_date': archive_date,
            'is_consolidated_tax': is_consolidated_tax,
            'credit_code_raw': credit_code_raw,
            'credit_codes': json.dumps(credit_codes, ensure_ascii=False),
            'contract_amount': None,
            'is_anchor': 0,
            'parent_contract_id': None,
            'is_payment': 1 if is_payment else 0,
            'source': LEGACY_SOURCE,
            'legacy_sheet': SHEET_NAME,
            'legacy_row': excel_row,
            'legacy_extra': json.dumps(legacy_extra, ensure_ascii=False) if legacy_extra['issues'] or legacy_extra['raw'] else '',
            'created_by': 'legacy-import',
            'created_by_name': '历史台账导入',
        })

    return {
        'records': records,
        'report_rows': report_rows,
        'issue_counts': issue_counts,
        'nature_values': nature_values,
        'payment_count': payment_count,
        'consolidated_count': consolidated_count,
        'duplicate_names': duplicate_names,
        'seq_min': min(used_seqs) if used_seqs else 0,
        'seq_max': max(used_seqs) if used_seqs else 0,
    }


def print_summary(result):
    records = result['records']
    print(f"有效合同行：{len(records)}")
    print(f"序号范围：{result['seq_min']} - {result['seq_max']}")
    print(f"付款类合同：{result['payment_count']}")
    print(f"汇总纳税：{result['consolidated_count']}")
    print(f"合同性质不同值：{len(result['nature_values'])}")
    print(f"异常行数：{len(result['report_rows'])}")
    if result['issue_counts']:
        print('异常字段统计：')
        for field, count in sorted(result['issue_counts'].items(), key=lambda x: (-x[1], x[0])):
            print(f'  {field}: {count}')
    duplicates = {k: v for k, v in result['duplicate_names'].items() if v > 1}
    if duplicates:
        print(f'重名合同：{len(duplicates)} 个名称')
        for name, count in list(duplicates.items())[:10]:
            print(f'  {name} x{count}')
    if result['report_rows']:
        print('\n异常样例：')
        for row in result['report_rows'][:12]:
            print(f"  行{row['行号']} 序号{row['序号']} {row['项目名称'][:30]} | {row['问题字段']} | {row['建议处理']}")


def write_report(path, rows):
    if not rows:
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['行号', '序号', '项目名称', '问题字段', '原始值', '建议处理'])
        writer.writeheader()
        writer.writerows(rows)
    print(f'异常报告：{path}')


def apply_import(db_path, records):
    if not os.path.exists(db_path):
        raise RuntimeError(f'数据库不存在：{db_path}')
    conn = sqlite3.connect(db_path)
    try:
        cols = {r[1] for r in conn.execute('PRAGMA table_info(contracts)').fetchall()}
        required = {
            'id', 'seq_no', 'name', 'counterparty', 'total_investment', 'nature_raw',
            'contract_content', 'project_manager', 'tech_manager', 'contract_date_raw',
            'contract_date', 'contract_date_end', 'agency_fee', 'stamp_tax_due',
            'stamp_tax_paid', 'copies', 'archive_date', 'is_consolidated_tax',
            'credit_code_raw', 'credit_codes', 'contract_amount', 'is_anchor',
            'parent_contract_id', 'is_payment', 'source', 'legacy_sheet',
            'legacy_row', 'legacy_extra', 'created_by', 'created_by_name',
        }
        missing = required - cols
        if missing:
            raise RuntimeError('contracts 表缺少字段：' + '、'.join(sorted(missing)))

        sql = '''
        INSERT INTO contracts (
          id, seq_no, name, counterparty, total_investment,
          nature_raw, contract_content, project_manager, tech_manager,
          contract_date_raw, contract_date, contract_date_end,
          agency_fee, stamp_tax_due, stamp_tax_paid, copies, archive_date,
          is_consolidated_tax, credit_code_raw, credit_codes, contract_amount,
          is_anchor, parent_contract_id, is_payment, source,
          legacy_sheet, legacy_row, legacy_extra, created_by, created_by_name
        ) VALUES (
          :id, :seq_no, :name, :counterparty, :total_investment,
          :nature_raw, :contract_content, :project_manager, :tech_manager,
          :contract_date_raw, :contract_date, :contract_date_end,
          :agency_fee, :stamp_tax_due, :stamp_tax_paid, :copies, :archive_date,
          :is_consolidated_tax, :credit_code_raw, :credit_codes, :contract_amount,
          :is_anchor, :parent_contract_id, :is_payment, :source,
          :legacy_sheet, :legacy_row, :legacy_extra, :created_by, :created_by_name
        )
        ON CONFLICT(legacy_sheet, legacy_row) DO UPDATE SET
          seq_no = excluded.seq_no,
          name = excluded.name,
          counterparty = excluded.counterparty,
          total_investment = excluded.total_investment,
          nature_raw = excluded.nature_raw,
          contract_content = excluded.contract_content,
          project_manager = excluded.project_manager,
          tech_manager = excluded.tech_manager,
          contract_date_raw = excluded.contract_date_raw,
          contract_date = excluded.contract_date,
          contract_date_end = excluded.contract_date_end,
          agency_fee = excluded.agency_fee,
          stamp_tax_due = excluded.stamp_tax_due,
          stamp_tax_paid = excluded.stamp_tax_paid,
          copies = excluded.copies,
          archive_date = excluded.archive_date,
          is_consolidated_tax = excluded.is_consolidated_tax,
          credit_code_raw = excluded.credit_code_raw,
          credit_codes = excluded.credit_codes,
          contract_amount = excluded.contract_amount,
          legacy_extra = excluded.legacy_extra,
          updated_at = datetime('now')
        '''
        conn.executemany(sql, records)
        conn.commit()
        print(f'已导入/更新 {len(records)} 条合同')
    finally:
        conn.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--xlsx', default=DEFAULT_XLSX)
    parser.add_argument('--sheet', default=SHEET_NAME)
    parser.add_argument('--db', default=DEFAULT_DB)
    parser.add_argument('--report', default=DEFAULT_REPORT)
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()

    if not os.path.exists(args.xlsx):
        print(f'Excel 文件不存在：{args.xlsx}', file=sys.stderr)
        sys.exit(1)

    result = build_records(args.xlsx, args.sheet)
    print_summary(result)
    write_report(args.report, result['report_rows'])

    if args.apply:
        apply_import(args.db, result['records'])
    elif not args.dry_run:
        print('\n当前为预览模式；确认无误后加 --apply 写入数据库。')


if __name__ == '__main__':
    main()